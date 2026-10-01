import unittest
from collections import OrderedDict
from pathlib import Path
import sys
import tempfile
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from excel_writer import ExcelWriter
from models import ParsedResult, StudentRecord, SubjectRecord
from parser_engine import ResultParsingEngine
from parser_factory import ParserFactory
from parsers.be_result_parser import BeStudentResultParser
from parsers.granite_fallback_parser import GraniteFallbackParser
from parsers.text_fallback_parser import TextFallbackParser
from parsers.ledger_parser import CODE, CollegeLedgerParser
from result_analyzer import ResultAnalyzer
from result_validator import assess_result
from utils import clean_result_value


class FeParsingTests(unittest.TestCase):
    def test_wrapped_mother_name_never_accepts_a_course_row(self) -> None:
        parser = CollegeLedgerParser()
        header = "PRN: 72413575K SEAT NO.: F500340001 NAME: KALASE PRANITA PRAKASH Mother's Name :-"
        student = StudentRecord(prn="72413575K", seat_no="F500340001")

        self.assertTrue(parser._needs_identity_continuation([header]))
        self.assertFalse(parser._is_identity_continuation([{"text": "AEC-101"}], "AEC-101 --- --- P 023 023 2 2 O 10 20"))

        parser._populate_identity(student, [header, "KUMUD"])
        self.assertEqual(student.name, "KALASE PRANITA PRAKASH")
        self.assertEqual(student.mother_name, "KUMUD")

    def test_all_dash_assessment_fields_are_excluded_before_excel_write(self) -> None:
        subject = SubjectRecord("AEC-101", "PROFESSIONAL COMMUNICATION SKILLS", OrderedDict({"CCE": "---", "TW": "023"}))
        student = StudentRecord(
            seat_no="F500340001",
            prn="72413575K",
            name="KALASE PRANITA PRAKASH",
            mother_name="KUMUD",
            semesters=OrderedDict({"Semester 1": OrderedDict({subject.label: subject})}),
        )
        result = ParsedResult(
            "College Ledger",
            "sample.pdf",
            [student],
            OrderedDict({"Semester 1": OrderedDict({subject.label: ["CCE", "TW"]})}),
        )

        dataframe = ExcelWriter.dataframe(result)
        self.assertNotIn(("Semester 1", subject.label, "CCE"), dataframe.columns)
        self.assertIn(("Semester 1", subject.label, "TW"), dataframe.columns)

    def test_missing_sgpa_is_calculated_in_excel_only_when_all_grades_pass(self) -> None:
        subject_labels = ["PCC-201 DATA STRUCTURES", "BSC-201 MATHEMATICS"]
        schema = OrderedDict({
            "Semester 1": OrderedDict(
                (label, ["Crd", "Grd", "GP"]) for label in subject_labels
            ),
        })

        def make_student(seat_no, first_grade="A", sgpa=None):
            summary = OrderedDict({"Credits Earned/Total": "5/5", "Total Credit Points": "43"})
            if sgpa is not None:
                summary["Semester 1 SGPA"] = sgpa
            subjects = OrderedDict({
                subject_labels[0]: SubjectRecord(
                    "PCC-201", "DATA STRUCTURES", OrderedDict({"Crd": "3", "Grd": first_grade, "GP": "9"})
                ),
                subject_labels[1]: SubjectRecord(
                    "BSC-201", "MATHEMATICS", OrderedDict({"Crd": "2", "Grd": "B", "GP": "8"})
                ),
            })
            return StudentRecord(
                seat_no=seat_no,
                semesters=OrderedDict({"Semester 1": subjects}),
                summary=summary,
            )

        result = ParsedResult(
            "College Ledger",
            "sample.pdf",
            [
                make_student("S1"),
                make_student("S2", first_grade="F"),
                make_student("S3", first_grade="FF"),
                make_student("S4", sgpa="7.25"),
            ],
            schema,
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "results.xlsx"
            ExcelWriter().write(result, workbook_path)
            worksheet = openpyxl.load_workbook(workbook_path, data_only=True)["Results"]
            sgpa_column = next(
                column for column in range(1, worksheet.max_column + 1)
                if worksheet.cell(3, column).value == "Semester 1 SGPA"
            )

            self.assertEqual(worksheet.cell(4, sgpa_column).value, 8.6)
            self.assertIsNone(worksheet.cell(5, sgpa_column).value)
            self.assertIsNone(worksheet.cell(6, sgpa_column).value)
            self.assertEqual(worksheet.cell(7, sgpa_column).value, "7.25")

            analyzed = ResultAnalyzer().analyze_excel(workbook_path)
            topper = next(student for student in analyzed["proforma_c"] if student["seat"] == "S1")
            self.assertEqual(topper["sgpa"], "8.6")

    def test_uppercase_ledger_headers_create_distinct_repeated_assessment_fields(self) -> None:
        labels = ["ISE", "ESE", "PR", "OR", "PR", "TW", "ISE", "ESE", "TW", "TOT", "CRD", "ERN", "GRD", "GRD", "CRD"]
        words = [
            {"text": label, "x0": index * 20, "x1": index * 20 + 10}
            for index, label in enumerate(labels)
        ]

        layout = CollegeLedgerParser._layout(words)
        self.assertIsNotNone(layout)
        self.assertEqual(
            [field for field, _ in layout.fields],
            ["ISE", "ESE", "PR", "OR", "PR (2)", "TW", "ISE (2)", "ESE (2)", "TW (2)", "Tot", "Result Status", "Crd", "Ern", "Grd", "GP", "CP"],
        )
        self.assertTrue(CODE.fullmatch("PCC-201-COM"))
        self.assertTrue(CODE.fullmatch("304191AB_6"))
        self.assertEqual(CollegeLedgerParser._clean_cell_value("028$"), "028")
        self.assertEqual(CollegeLedgerParser._clean_cell_value("028#"), "028")
        self.assertEqual(CollegeLedgerParser._clean_cell_value("028$&"), "028")

    def test_available_fe_pdfs_capture_sgpa_summary(self) -> None:
        root = Path(__file__).resolve().parents[1] / "pdfs"
        engine = ResultParsingEngine()
        pdfs = [root / name for name in ["FE 2024 pattern.pdf", "FE 2025.pdf"] if (root / name).exists()]
        if not pdfs:
            self.skipTest("No FE sample PDFs are present in pdfs/")

        for pdf_path in pdfs:
            result = engine.parse(pdf_path)
            self.assertGreater(len(result.students), 0, pdf_path.name)
            student = result.students[0]
            self.assertTrue(
                any("SGPA" in field for field in student.summary),
                f"{pdf_path.name} should contain an SGPA summary",
            )

    def test_available_uppercase_ledger_pdfs_detect_subjects(self) -> None:
        root = Path(__file__).resolve().parents[1] / "pdfs"
        names = [
            "SE Comp 2024 pattern summer 2026 result.pdf",
            "SE E&TC 2024 pattern Summer 2026 result.pdf",
            "TE Computer 2019 Pattern.pdf",
            "TE E&TC 2019 May 2026 (1).pdf",
        ]
        pdfs = [root / name for name in names if (root / name).exists()]
        if not pdfs:
            self.skipTest("No uppercase-header ledger samples are present in pdfs/")

        for pdf_path in pdfs:
            result = ResultParsingEngine().parse(pdf_path)
            self.assertGreater(len(result.students), 0, pdf_path.name)
            self.assertGreater(
                sum(len(subjects) for subjects in result.schema.values()),
                0,
                f"{pdf_path.name} should contain detected subjects",
            )
            self.assertGreater(
                sum(len(subjects) for subjects in result.students[0].semesters.values()),
                0,
                f"{pdf_path.name} should contain parsed student subjects",
            )
            self.assertFalse(result.requires_review, f"{pdf_path.name} should pass the quality gate")

    def test_invalid_mark_character_requires_review(self) -> None:
        subject = SubjectRecord("PCC-201", "Sample", OrderedDict({"Tot": "$45"}))
        student = StudentRecord(
            seat_no="S1",
            semesters=OrderedDict({"Semester 1": OrderedDict({subject.label: subject})}),
        )
        result = assess_result(
            ParsedResult(
                "College Ledger",
                "bad.pdf",
                [student],
                OrderedDict({"Semester 1": OrderedDict({subject.label: ["Tot"]})}),
            )
        )
        self.assertTrue(result.requires_review)
        self.assertTrue(any("invalid marks character" in note for note in result.review_notes))

    def test_standalone_status_marker_never_creates_a_merged_mark(self) -> None:
        header = ["ISE", "ESE", "TOT", "CRD", "ERN", "GRD", "GRD", "CRD"]
        layout = CollegeLedgerParser._layout(
            [{"text": label, "x0": index * 20, "x1": index * 20 + 10} for index, label in enumerate(header)]
        )
        assert layout is not None
        subject = CollegeLedgerParser._subject(
            [
                {"text": "PCC-201", "x0": -40, "x1": -10},
                {"text": "*", "x0": 0, "x1": 2},
                {"text": "027", "x0": 0, "x1": 10},
                {"text": "*", "x0": 20, "x1": 22},
                {"text": "047", "x0": 20, "x1": 30},
                {"text": "074", "x0": 40, "x1": 50},
            ],
            layout,
            {},
        )
        assert subject is not None
        self.assertEqual(subject.fields["ISE"], "027")
        self.assertEqual(subject.fields["ESE"], "047")
        self.assertFalse(any(" | " in value for value in subject.fields.values()))

    def test_numeric_footnote_suffixes_are_not_exported_as_marks(self) -> None:
        self.assertEqual(clean_result_value("10$/025"), "10/025")
        self.assertEqual(clean_result_value("2$"), "2")
        self.assertEqual(clean_result_value("028$&"), "028")
        self.assertEqual(clean_result_value("A+"), "A+")

    def test_granite_recovery_accepts_only_source_grounded_records(self) -> None:
        source = "PRN: 123 SEAT NO.: S1 NAME: A TEST\nPCC-201 DATA STRUCTURES 20 30 50"
        payload = {
            "records": [{
                "prn": "123",
                "seat_no": "S1",
                "name": "A TEST",
                "semesters": [{
                    "label": "Semester 1",
                    "subjects": [{"code": "PCC-201", "name": "DATA STRUCTURES", "fields": {"ISE": "20", "ESE": "30", "Tot": "50"}}],
                }],
                "summary": {},
            }],
        }
        result = GraniteFallbackParser._from_payload(payload, "unknown.pdf", source)
        self.assertIsNotNone(result)
        assert result is not None
        self.assertEqual(result.students[0].semesters["Semester 1"]["PCC-201 DATA STRUCTURES"].fields["Tot"], "50")
        self.assertTrue(result.requires_review)

        payload["records"][0]["prn"] = "invented"
        self.assertIsNone(GraniteFallbackParser._from_payload(payload, "unknown.pdf", source))

    def test_failed_known_parser_never_hands_off_to_granite(self) -> None:
        class BrokenParser:
            pdf_type = "College Ledger"

            @staticmethod
            def parse(_pdf_path):
                raise ValueError("different table")

        subject = SubjectRecord("PCC-201", "Sample", OrderedDict({"Tot": "50"}))
        recovered = ParsedResult(
            "Text Fallback",
            "unknown.pdf",
            [StudentRecord(seat_no="S1", semesters=OrderedDict({"Semester 1": OrderedDict({subject.label: subject})}))],
            OrderedDict({"Semester 1": OrderedDict({subject.label: ["Tot"]})}),
            raw_text="S1 PCC-201 50",
            requires_review=True,
        )
        with patch("parser_engine.ParserFactory.for_pdf", return_value=BrokenParser()):
            with patch("parser_engine.TextFallbackParser.parse", return_value=recovered) as recovery:
                with patch("parser_engine.GraniteFallbackParser.parse") as granite:
                    result = ResultParsingEngine().parse(Path("unknown.pdf"))
        self.assertTrue(result.requires_review)
        self.assertEqual(recovery.call_count, 1)
        self.assertEqual(granite.call_count, 0)
        self.assertTrue(any("source text was preserved" in note for note in result.review_notes))

    def test_available_2019_student_table_is_never_routed_to_granite(self) -> None:
        pdf_path = Path(__file__).resolve().parents[1] / "pdfs" / "TE E&TC 2019 (Second) May 2026 (4).pdf"
        if not pdf_path.exists():
            self.skipTest("No 2019 student-table regression PDF is present in pdfs/")

        self.assertIsInstance(ParserFactory.for_pdf(pdf_path), BeStudentResultParser)
        result = ResultParsingEngine().parse(pdf_path)
        self.assertEqual(result.pdf_type, "Student Result")
        self.assertFalse(result.requires_review)
        self.assertEqual(len(result.students), 28)
        self.assertGreater(sum(len(subjects) for subjects in result.schema.values()), 0)


if __name__ == "__main__":
    unittest.main()
